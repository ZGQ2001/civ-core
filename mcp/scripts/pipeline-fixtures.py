"""跨语言冒烟的测试夹具；生成最小 Word 模板、填入涂层样例并检查产物。"""

from __future__ import annotations

import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

root = Path(sys.argv[2])
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

if sys.argv[1] == "prepare":
    # 使用已生成的官方输入模板；25mm 为明确的合格测试样例，并非现场实测。
    wb = load_workbook(root / "coating-input.xlsx")
    for ws in wb:
        if ws.title.startswith("测点数据"):
            headers = [c.value for c in ws[1]]
            section_col = next(i + 1 for i, h in enumerate(headers) if h in ("截面号", "处号"))
            for row in ws.iter_rows(min_row=2, min_col=section_col + 1):
                for cell in row:
                    cell.value = 25
    wb.save(root / "coating-input.xlsx")
    wb.close()

    paragraphs = [
        "自动化验收样例（非现场检测报告）",
        "工程名称：{{工程名称}}",
        "{{表格:锚杆}}",
        "{{表格:防火涂层}}",
    ]
    document = ET.Element(f"{{{W}}}document")
    body = ET.SubElement(document, f"{{{W}}}body")
    for text in paragraphs:
        p = ET.SubElement(body, f"{{{W}}}p")
        r = ET.SubElement(p, f"{{{W}}}r")
        ET.SubElement(r, f"{{{W}}}t").text = text
    section = ET.SubElement(body, f"{{{W}}}sectPr")
    ET.SubElement(section, f"{{{W}}}pgSz", {f"{{{W}}}w": "11906", f"{{{W}}}h": "16838"})
    ET.SubElement(
        section,
        f"{{{W}}}pgMar",
        {f"{{{W}}}{k}": "1134" for k in ("top", "bottom", "left", "right")},
    )
    with ZipFile(root / "template.docx", "w", ZIP_DEFLATED) as z:
        z.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        z.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        )
        z.writestr(
            "word/document.xml", ET.tostring(document, encoding="utf-8", xml_declaration=True)
        )
elif sys.argv[1] == "verify":
    wb = load_workbook(root / "anchor-result.xlsx", data_only=True)
    assert "判定依据" in " ".join(wb.sheetnames), wb.sheetnames
    assert "_批次参数" in wb.sheetnames
    assert wb["_批次参数"].sheet_state == "hidden"
    values = [c.value for row in wb["批次1-数据分析"] for c in row]
    # 自带锚杆样例：2.63 - 0.58 = 2.05mm 弹性位移。
    assert any(isinstance(v, (int, float)) and abs(v - 2.05) < 1e-9 for v in values)
    wb.close()
    with ZipFile(root / "combined-report.docx") as z:
        doc = ET.fromstring(z.read("word/document.xml"))
        text = "".join(doc.itertext())
        assert "自动化验收工程" in text
        assert "{{" not in text, text
        assert len(doc.findall(f".//{{{W}}}tbl")) == 5
        assert len(doc.findall(f".//{{{W}}}drawing")) == 3
        assert len([n for n in z.namelist() if n.lower().endswith(".png")]) == 3
    print("Verified: Excel traceability, displacement 2.05mm, 5 Word tables, 3 embedded curves.")
else:
    raise ValueError(f"Unknown mode: {sys.argv[1]}")
